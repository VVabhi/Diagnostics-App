from app import app, employees_collection, patient_collection, inventory_collection, store_inventory_collection, LoginForm, RegistrationForm
from flask import render_template, request, redirect, url_for, session, jsonify, send_file
from bson import ObjectId
import datetime
from docx import Document
from io import BytesIO


@app.route('/')
def index():
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    form = LoginForm()
    if form.validate_on_submit():
        username = form.username.data
        password = form.password.data

        employee = employees_collection.find_one({'username': username, 'password': password})

        if employee:
            session['username'] = username
            return redirect(url_for('dashboard'))
        else:
            return render_template('login.html', form=form, error="Invalid credentials")

    return render_template('login.html', form=form)

@app.route('/register', methods=['GET', 'POST'])
def register():
    form = RegistrationForm()
    if form.validate_on_submit():
        employee_name = form.employee_name.data
        employee_id = form.employee_id.data
        username = form.username.data
        password = form.password.data

        employee_data = {
            'employee_name': employee_name,
            'employee_id': employee_id,
            'username': username,
            'password': password,
        }

        employees_collection.insert_one(employee_data)
        return redirect(url_for('login'))

    return render_template('register.html', form=form)

@app.route('/dashboard')
def dashboard():
    if 'username' in session:
        username = session['username']
        employee = employees_collection.find_one({'username': username})
        return render_template('dashboard.html', employee=employee)
    else:
        return redirect(url_for('login'))


@app.route('/input')
def input_page():
    if 'username' in session:
        return render_template('input.html')
    else:
        return redirect(url_for('login'))

@app.route('/submit', methods=['POST'])
def submit():
    data = request.get_json()
    patientName = data.get('patientName')
    phoneNumber = data.get('phoneNumber')
    age = data.get('age')
    sex = data.get('sex')
    refBy = data.get('refBy')
    date = data.get('date')
    test = data.get('test')
    otherTest = data.get('otherTest')
    cost = data.get('cost')
    modeOfPayment = data.get('modeOfPayment')

    patient_data = {
        'patient_name': patientName,
        'phone_number': phoneNumber,
        'age': age,
        'sex': sex,
        'ref_by': refBy,
        'date': date,
        'test': test,
        'other_test': otherTest,
        'cost': cost,
        'mode_of_payment': modeOfPayment
    }

    patient_collection.insert_one(patient_data)
    return jsonify({"message": "Patient details added successfully!"}), 201

@app.route('/display')
def display():
    if 'username' in session:
        patients_list = list(patient_collection.find({'report': {'$exists': False}}))  # Filter out patients with reports
        return render_template('test.html', patients=patients_list)
    else:
        return redirect(url_for('login'))

@app.route('/generate_word_report/<patient_id>')
def generate_word_report(patient_id):
    # Fetch patient details from the database using patient_id
    patient = patient_collection.find_one({'_id': ObjectId(patient_id)})
    
    if not patient:
        return "Patient not found", 404

    # Create a Word document
    doc = Document()
    doc.add_heading('Patient Report', level=1)
    doc.add_paragraph(f"Name: {patient.get('patient_name', 'N/A')}")
    doc.add_paragraph(f"Age: {patient.get('age', 'N/A')}")
    doc.add_paragraph(f"Sex: {patient.get('sex', 'N/A')}")
    doc.add_paragraph(f"Referred By: {patient.get('ref_by', 'N/A')}")
    doc.add_paragraph(f"Date: {patient.get('date', 'N/A')}")
    doc.add_paragraph(f"Test: {patient.get('test', 'N/A')}")
    if patient.get('other_test'):
        doc.add_paragraph(f"Other Test: {patient.get('other_test')}")

    # Save the document to a BytesIO object
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return send_file(
        file_stream,
        as_attachment=True,
        download_name=f"patient_report_{patient_id}.docx",
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

@app.route('/generate_report', methods=['POST'])
def generate_report():
    data = request.get_json()
    patient_id = data.get('patient_id')
    report = data.get('report')

    # Update patient record in MongoDB with the report
    patient_collection.update_one({'_id': ObjectId(patient_id)}, {'$set': {'report': report}})

    return jsonify({"message": "Report updated successfully!"}), 200

@app.route('/reports')
def reports_page():
    if 'username' in session:
        patients_list = list(patient_collection.find({'report': {'$exists': True}}))
        return render_template('reports.html', patients=patients_list)
    else:
        return redirect(url_for('login'))

@app.route('/payments')
def payments_page():
    if 'username' in session:
        date_filter = request.args.get('date')
        query = {}
        if date_filter:
            query['date'] = date_filter

        payments_list = list(patient_collection.find(query, {'patient_name': 1, 'cost': 1, 'mode_of_payment': 1, 'date': 1}))

        # Convert ObjectId to string
        for payment in payments_list:
            payment['_id'] = str(payment['_id'])

        total_costs = {'cash': 0, 'upi': 0, 'card': 0}
        for payment in payments_list:
            mode_of_payment = payment.get('mode_of_payment', '').lower()
            if mode_of_payment in total_costs:
                total_costs[mode_of_payment] += float(payment.get('cost', 0))

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({
                'payments': payments_list,
                'total_costs': total_costs
            })

        return render_template('payments.html', payments=payments_list, total_costs=total_costs)
    else:
        return redirect(url_for('login'))

@app.route('/inventory', methods=['GET', 'POST'])
def inventory_page():
    if 'username' in session:
        if request.method == 'POST':
            if request.is_json:
                inventory_list = list(inventory_collection.find({}, {'_id': 0}))  
                return jsonify(inventory_list)
            else:
                name = request.form['name']
                apperatus = request.form['apperatus']
                quantity = request.form['quantity']
                supplier = request.form['supplier']
                cost = request.form['cost']
                payment_date = request.form['payment_date']

                inventory_data = {
                    'name': name,
                    'apperatus': apperatus,
                    'quantity': quantity,
                    'supplier': supplier,
                    'cost': cost,
                    'payment_date': payment_date
                }

                inventory_collection.insert_one(inventory_data)
                return redirect(url_for('inventory_page'))

        inventory_list = list(inventory_collection.find())
        return render_template('inventory.html', inventory=inventory_list)
    else:
        return redirect(url_for('login'))

@app.route('/store_inventory', methods=['GET', 'POST'])
def store_inventory_page():
    if 'username' in session:
        if request.method == 'POST':
            apparatus_name = request.form['apparatusName']
            quantity = request.form['quantity']

            # Separate the numeric part and the unit part of the quantity
            quantity_value = ''.join(filter(str.isdigit, quantity))
            quantity_unit = ''.join(filter(str.isalpha, quantity))

            if not quantity_value or not quantity_unit:
                return "Invalid quantity input. Please enter a valid quantity with units (e.g., 25L, 20Kg).", 400

            try:
                quantity_value = int(quantity_value)

                # Check if the apparatus already exists in the collection
                existing_item = store_inventory_collection.find_one({'apparatus_name': apparatus_name})

                if existing_item:
                    # If it exists, update the quantity
                    existing_quantity_value = int(''.join(filter(str.isdigit, existing_item['quantity'])))
                    updated_quantity_value = existing_quantity_value + quantity_value
                    updated_quantity = f"{updated_quantity_value}{quantity_unit}"
                    store_inventory_collection.update_one(
                        {'_id': existing_item['_id']},
                        {'$set': {'quantity': updated_quantity}}
                    )
                else:
                    # If it doesn't exist, insert a new document
                    store_inventory_data = {
                        'apparatus_name': apparatus_name,
                        'quantity': quantity
                    }
                    store_inventory_collection.insert_one(store_inventory_data)

                return redirect(url_for('store_inventory_page'))
            except ValueError:
                # Handle invalid quantity input
                return "Invalid quantity input. Please enter a valid number.", 400

        inventory_list = list(store_inventory_collection.find())
        return render_template('store_inventory.html', inventory=inventory_list)
    else:
        return redirect(url_for('login'))

@app.route('/logout')
def logout():
    session.pop('username', None)
    return redirect(url_for('login'))


